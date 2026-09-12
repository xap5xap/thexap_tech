from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[5]
OUTPUT = Path("/tmp/xavier-perez-innodata-idi037-resume-draft.docx")

BLACK = "202124"
GRAY = "55585E"
LIGHT_GRAY = "73767C"
ORANGE = "F59415"
FONT = "Liberation Sans"


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float, color: str = BLACK, bold: bool = False, italic: bool = False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text: str, url: str, color: str = GRAY):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), "18")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, size_el, underline])
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_section_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=5, after=2, line=1.0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    set_run_font(run, 9.2, BLACK, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), ORANGE)
    borders.append(bottom)
    p_pr.append(borders)


def add_body_paragraph(doc: Document, text: str, size=9.25, color=BLACK, after=2.5):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.04)
    run = p.add_run(text)
    set_run_font(run, size, color)
    return p


def add_bullet(doc: Document, text: str, after=1.3):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.03)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.keep_together = True
    bullet = p.add_run("•")
    set_run_font(bullet, 9.1, ORANGE, bold=True)
    spacer = p.add_run("  ")
    set_run_font(spacer, 9.1)
    body = p.add_run(text)
    set_run_font(body, 9.1)
    return p


def add_role(doc: Document, role: str, organization: str, dates: str):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.85)
    table.columns[1].width = Inches(1.15)
    remove_table_borders(table)
    row = table.rows[0]
    set_repeat_table_header(row)
    for cell in row.cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    left = row.cells[0].paragraphs[0]
    set_paragraph_spacing(left, after=0, line=1.0)
    left.paragraph_format.keep_with_next = True
    r = left.add_run(role)
    set_run_font(r, 9.5, BLACK, bold=True)
    r = left.add_run(f" | {organization}")
    set_run_font(r, 9.35, GRAY)
    right = row.cells[1].paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(right, after=0, line=1.0)
    right.paragraph_format.keep_with_next = True
    r = right.add_run(dates)
    set_run_font(r, 8.8, LIGHT_GRAY, bold=True)


def build_resume():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.43)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.25)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    doc.core_properties.title = "Xavier Perez Resume"
    doc.core_properties.subject = "Resume tailored for Innodata IDI037 SMEs for AI Evaluations"
    doc.core_properties.keywords = "full-stack, cloud, AWS, technical review, AI-assisted engineering"

    name = doc.add_paragraph()
    set_paragraph_spacing(name, after=0, line=0.95)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.add_run("XAVIER PEREZ")
    set_run_font(r, 24, BLACK, bold=True)

    title = doc.add_paragraph()
    set_paragraph_spacing(title, after=2.5, line=1.0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Senior Full-Stack Product Engineer | Software, Cloud, and Technical Evaluation")
    set_run_font(r, 10.5, ORANGE, bold=True)

    contact = doc.add_paragraph()
    set_paragraph_spacing(contact, after=0.5, line=1.0)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = contact.add_run("Quito, Ecuador  |  Spanish: Native  |  English: Fluent")
    set_run_font(r, 9.0, GRAY)

    links = doc.add_paragraph()
    set_paragraph_spacing(links, after=2, line=1.0)
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(links, "thexap.com", "https://www.thexap.com/")
    r = links.add_run("  |  ")
    set_run_font(r, 9.0, LIGHT_GRAY)
    add_hyperlink(links, "Upwork profile", "https://www.upwork.com/freelancers/~01c4dd86006663beac")
    r = links.add_run("  |  ")
    set_run_font(r, 9.0, LIGHT_GRAY)
    add_hyperlink(links, "GitHub", "https://github.com/xap5xap")

    add_section_heading(doc, "Professional Summary")
    add_body_paragraph(
        doc,
        "Senior full-stack product engineer with 9+ years delivering production software across web applications, "
        "backend services, APIs, databases, testing, and cloud infrastructure. Experienced in turning requirements "
        "into architecture, reviewing technical work, identifying risks, and verifying results. Uses AI-assisted "
        "workflows for research, implementation, and review while remaining responsible for final decisions.",
        after=1.5,
    )

    add_section_heading(doc, "Relevant Expertise")
    expertise = doc.add_table(rows=2, cols=3)
    expertise.autofit = False
    remove_table_borders(expertise)
    set_repeat_table_header(expertise.rows[0])
    widths = [Inches(2.38), Inches(2.38), Inches(2.38)]
    items = [
        "Architecture and code review",
        "AWS and cloud integrations",
        "REST and GraphQL APIs",
        "Authentication and multi-tenancy",
        "Testing and failure analysis",
        "Technical documentation",
    ]
    for index, cell in enumerate([cell for row in expertise.rows for cell in row.cells]):
        cell.width = widths[index % 3]
        set_cell_margins(cell, top=35, start=70, bottom=35, end=70)
        shade_cell(cell, "F4F5F6")
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(items[index])
        set_run_font(r, 8.7, BLACK, bold=True)

    add_section_heading(doc, "Experience")
    add_role(doc, "Independent Senior Full-Stack Product Engineer", "Upwork and direct clients", "2016 to present")
    add_bullet(
        doc,
        "Completed 28 Upwork jobs and 15,309 hours with 100% Job Success and Top Rated Plus status, as captured in September 2026.",
    )
    add_bullet(
        doc,
        "Delivered and maintained React, Next.js, TypeScript, Node.js, mobile, and web applications in new and established codebases for distributed teams.",
    )
    add_bullet(
        doc,
        "Worked with AWS Cognito, API Gateway, Lambda, and CDK for authentication, APIs, serverless compute, and infrastructure as code.",
    )
    add_bullet(
        doc,
        "Defined requirements, analyzed trade-offs, implemented changes, investigated failures, and documented handoffs across product and engineering work.",
    )
    add_bullet(
        doc,
        "Built and reviewed unit, integration, Cypress, and end-to-end tests to verify behavior against expected outcomes.",
        after=2.0,
    )

    add_role(doc, "Product and Engineering Lead", "Armonía", "2023 to 2026")
    add_bullet(
        doc,
        "Led discovery, UX, architecture, and full-stack delivery of a public website and private operating application for a psychology practice.",
    )
    add_bullet(
        doc,
        "Replaced an expanding operational spreadsheet through a staged transition; the practice owner now uses the dedicated application as her primary system.",
    )
    add_bullet(
        doc,
        "Implemented React, TypeScript, Node.js/Hono, PostgreSQL, Clerk, OAuth, webhooks, and Google Calendar synchronization with clear privacy boundaries.",
        after=2.0,
    )

    add_section_heading(doc, "AI and Evaluation Approach")
    add_bullet(
        doc,
        "Uses prompts and AI coding tools to support research, implementation, and technical review, then checks conclusions against source material and acceptance criteria.",
    )
    add_bullet(
        doc,
        "Separates verified results from assumptions, records limitations, and gives precise feedback on correctness, risk, maintainability, and missing evidence.",
        after=1.5,
    )

    add_section_heading(doc, "Certification")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=0, line=1.0)
    r = p.add_run("AWS Certified AI Practitioner")
    set_run_font(r, 9.3, BLACK, bold=True)
    r = p.add_run("  |  Issued April 2026  |  Expires April 2029")
    set_run_font(r, 9.1, GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(footer, after=0, line=1.0)
    r = footer.add_run("Xavier Perez | Resume | September 2026")
    set_run_font(r, 7.5, LIGHT_GRAY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
